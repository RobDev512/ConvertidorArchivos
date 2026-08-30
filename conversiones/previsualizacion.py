import io
import os
import shutil
import subprocess
import tempfile

from PIL import Image, ImageOps

from funciones.formatos import (
    obtener_categoria_archivo,
    obtener_formato_archivo
)

from funciones.capacidades import (
    detectar_imagemagick,
    detectar_microsoft_powerpoint
)

from conversiones.documentos_texto import convertir_documento_texto
from conversiones.presentaciones import convertir_presentacion


# ----------------------------------- #
#      Límites de previsualización    #
# ----------------------------------- #

MAXIMO_FILAS_EXCEL = 300
MAXIMO_COLUMNAS_EXCEL = 80

MAXIMO_IMAGEN_PREVISUALIZACION = 2600

DPI_PREVISUALIZACION_PDF = 140


# ================================================================
# OBTENER TIPO DE PREVISUALIZACIÓN
# ================================================================
def obtener_tipo_previsualizacion(ruta_archivo):

    formato = obtener_formato_archivo(
        ruta_archivo
    )

    categoria = obtener_categoria_archivo(
        ruta_archivo
    )

    if categoria == "imagen":
        return "imagen"

    if formato in {
        "PDF",
        "DOCX"
    }:
        return "documento"

    if formato == "PPTX":
        return "presentacion"

    if formato == "XLSX":
        return "hoja_calculo"

    return "archivo"


# ================================================================
# CREAR OPCIONES DE SUBPROCESO BINARIO
# ================================================================
def crear_opciones_subproceso_binario():

    opciones = {
        "stdout": subprocess.PIPE,
        "stderr": subprocess.PIPE
    }

    if os.name == "nt":
        opciones["creationflags"] = subprocess.CREATE_NO_WINDOW

    return opciones


# ================================================================
# CONVERTIR IMAGEN A PNG CON PILLOW
# ================================================================
def convertir_imagen_pillow_a_png(
    ruta_archivo
):

    with Image.open(
        ruta_archivo
    ) as imagen:

        # Primer fotograma en imágenes animadas
        try:
            imagen.seek(0)

        except EOFError:
            pass

        # Respetar orientación EXIF
        imagen = ImageOps.exif_transpose(
            imagen
        )

        imagen = imagen.copy()

        # Mantener transparencia cuando exista
        if imagen.mode in {
            "RGBA",
            "LA"
        }:

            imagen = imagen.convert(
                "RGBA"
            )

        elif (
            imagen.mode == "P"
            and "transparency" in imagen.info
        ):

            imagen = imagen.convert(
                "RGBA"
            )

        elif imagen.mode not in {
            "RGB",
            "L"
        }:

            imagen = imagen.convert(
                "RGB"
            )

        # Evitar cargar imágenes gigantes a tamaño completo
        imagen.thumbnail(
            (
                MAXIMO_IMAGEN_PREVISUALIZACION,
                MAXIMO_IMAGEN_PREVISUALIZACION
            ),
            Image.Resampling.LANCZOS
        )

        buffer = io.BytesIO()

        imagen.save(
            buffer,
            format="PNG"
        )

        return buffer.getvalue()


# ================================================================
# CONVERTIR IMAGEN A PNG CON IMAGEMAGICK
# ================================================================
def convertir_imagen_imagemagick_a_png(
    ruta_archivo
):

    ruta_imagemagick = detectar_imagemagick()

    if not ruta_imagemagick:

        raise RuntimeError(
            "ImageMagick no está disponible."
        )

    comando = [
        ruta_imagemagick,
        f"{os.path.abspath(ruta_archivo)}[0]",
        "-auto-orient",
        "-thumbnail",
        (
            f"{MAXIMO_IMAGEN_PREVISUALIZACION}"
            f"x{MAXIMO_IMAGEN_PREVISUALIZACION}>"
        ),
        "png:-"
    ]

    resultado = subprocess.run(
        comando,
        **crear_opciones_subproceso_binario()
    )

    if resultado.returncode != 0:

        detalle = resultado.stderr.decode(
            "utf-8",
            errors="replace"
        ).strip()

        raise RuntimeError(
            "ImageMagick no pudo generar la previsualización."
            + (
                f"\n\n{detalle}"
                if detalle
                else ""
            )
        )

    if not resultado.stdout:

        raise RuntimeError(
            "ImageMagick no generó datos de imagen."
        )

    return resultado.stdout


# ================================================================
# PREPARAR PREVISUALIZACIÓN DE IMAGEN
# ================================================================
def preparar_previsualizacion_imagen(
    ruta_archivo
):

    error_pillow = None

    try:

        datos = convertir_imagen_pillow_a_png(
            ruta_archivo
        )

        return {
            "tipo": "imagen",
            "ruta_archivo": ruta_archivo,
            "datos_imagen": datos,
            "carpeta_temporal": None
        }

    except Exception as error:

        error_pillow = error


    # ----------------------------------- #
    #       Respaldo con ImageMagick      #
    # ----------------------------------- #

    try:

        datos = convertir_imagen_imagemagick_a_png(
            ruta_archivo
        )

        return {
            "tipo": "imagen",
            "ruta_archivo": ruta_archivo,
            "datos_imagen": datos,
            "carpeta_temporal": None
        }

    except Exception as error_imagemagick:

        raise RuntimeError(
            "No fue posible generar la previsualización "
            "de esta imagen.\n\n"
            f"Pillow: {error_pillow}\n\n"
            f"ImageMagick: {error_imagemagick}"
        ) from error_imagemagick


# ================================================================
# OBTENER CANTIDAD DE PÁGINAS PDF
# ================================================================
def obtener_cantidad_paginas_pdf(
    ruta_pdf
):

    import pypdfium2 as pdfium

    documento = pdfium.PdfDocument(
        ruta_pdf
    )

    try:

        cantidad = len(
            documento
        )

    finally:

        documento.close()

    if cantidad <= 0:

        raise RuntimeError(
            "El documento no contiene páginas."
        )

    return cantidad


# ================================================================
# RENDERIZAR PÁGINA PDF
# ================================================================
def renderizar_pagina_pdf(
    ruta_pdf,
    indice_pagina
):

    import pypdfium2 as pdfium

    documento = pdfium.PdfDocument(
        ruta_pdf
    )

    try:

        cantidad_paginas = len(
            documento
        )

        if (
            indice_pagina < 0
            or indice_pagina >= cantidad_paginas
        ):

            raise IndexError(
                "La página solicitada no existe."
            )

        pagina = documento[
            indice_pagina
        ]

        try:

            escala = (
                DPI_PREVISUALIZACION_PDF
                / 72
            )

            bitmap = pagina.render(
                scale=escala,
                rev_byteorder=True,
                prefer_bgrx=True
            )

            try:

                imagen = bitmap.to_pil()

                # Copia independiente de PDFium
                imagen = imagen.convert(
                    "RGB"
                )

                buffer = io.BytesIO()

                imagen.save(
                    buffer,
                    format="PNG"
                )

                return buffer.getvalue()

            finally:

                bitmap.close()

        finally:

            pagina.close()

    finally:

        documento.close()


# ================================================================
# PREPARAR PDF O DOCX
# ================================================================
def preparar_previsualizacion_documento(
    ruta_archivo
):

    formato = obtener_formato_archivo(
        ruta_archivo
    )

    carpeta_temporal = None


    # ----------------------------------- #
    #               PDF                   #
    # ----------------------------------- #

    if formato == "PDF":

        ruta_pdf = os.path.abspath(
            ruta_archivo
        )


    # ----------------------------------- #
    #              DOCX                   #
    # ----------------------------------- #

    elif formato == "DOCX":

        carpeta_temporal = tempfile.mkdtemp(
            prefix="convertidor_docx_"
        )

        ruta_pdf = os.path.join(
            carpeta_temporal,
            "previsualizacion.pdf"
        )

        try:

            convertir_documento_texto(
                ruta_archivo,
                ruta_pdf,
                "PDF"
            )

        except Exception:

            shutil.rmtree(
                carpeta_temporal,
                ignore_errors=True
            )

            raise


    else:

        raise RuntimeError(
            f"El formato {formato} no utiliza "
            "este previsualizador."
        )


    try:

        cantidad_paginas = obtener_cantidad_paginas_pdf(
            ruta_pdf
        )

        datos_imagen = renderizar_pagina_pdf(
            ruta_pdf,
            0
        )


        return {
            "tipo": "documento",
            "ruta_archivo": ruta_archivo,
            "ruta_pdf": ruta_pdf,
            "pagina_actual": 0,
            "cantidad_paginas": cantidad_paginas,
            "datos_imagen": datos_imagen,
            "carpeta_temporal": carpeta_temporal
        }


    except Exception:

        if carpeta_temporal:

            shutil.rmtree(
                carpeta_temporal,
                ignore_errors=True
            )

        raise


# ================================================================
# EXPORTAR DIAPOSITIVAS CON POWERPOINT
# ================================================================
def exportar_diapositivas_powerpoint(
    ruta_archivo,
    carpeta_destino
):

    try:

        import pythoncom
        import win32com.client

    except ImportError as error:

        raise RuntimeError(
            "pywin32 no está disponible."
        ) from error


    pythoncom.CoInitialize()

    aplicacion = None
    presentacion = None


    try:

        aplicacion = win32com.client.DispatchEx(
            "PowerPoint.Application"
        )

        try:
            aplicacion.DisplayAlerts = 1

        except Exception:
            pass

        try:
            aplicacion.AutomationSecurity = 3

        except Exception:
            pass


        # ----------------------------------- #
        #        Abrir presentación           #
        # ----------------------------------- #

        presentacion = aplicacion.Presentations.Open(
            os.path.abspath(
                ruta_archivo
            ),
            True,
            False,
            False
        )


        cantidad = presentacion.Slides.Count

        if cantidad <= 0:

            raise RuntimeError(
                "La presentación no contiene diapositivas."
            )


        # ----------------------------------- #
        #       Calcular resolución           #
        # ----------------------------------- #

        ancho_original = float(
            presentacion.PageSetup.SlideWidth
        )

        alto_original = float(
            presentacion.PageSetup.SlideHeight
        )

        ancho_salida = 1600

        alto_salida = round(
            ancho_salida
            * alto_original
            / ancho_original
        )


        # ----------------------------------- #
        #       Exportar diapositivas         #
        # ----------------------------------- #

        rutas = []

        cantidad_digitos = max(
            3,
            len(str(cantidad))
        )

        for indice in range(
            1,
            cantidad + 1
        ):

            ruta_imagen = os.path.join(
                carpeta_destino,
                (
                    f"diapositiva_"
                    f"{indice:0{cantidad_digitos}d}.png"
                )
            )

            diapositiva = presentacion.Slides(
                indice
            )

            diapositiva.Export(
                ruta_imagen,
                "PNG",
                ancho_salida,
                alto_salida
            )

            if not os.path.isfile(
                ruta_imagen
            ):

                raise RuntimeError(
                    f"No se pudo exportar la diapositiva {indice}."
                )

            rutas.append(
                ruta_imagen
            )


        return rutas


    finally:

        if presentacion is not None:

            try:
                presentacion.Close()

            except Exception:
                pass

        if aplicacion is not None:

            try:
                aplicacion.Quit()

            except Exception:
                pass

        pythoncom.CoUninitialize()


# ================================================================
# LEER IMAGEN DESDE ARCHIVO
# ================================================================
def leer_archivo_binario(
    ruta_archivo
):

    with open(
        ruta_archivo,
        "rb"
    ) as archivo:

        return archivo.read()


# ================================================================
# PREPARAR PPTX MEDIANTE PDF
# ================================================================
def preparar_pptx_mediante_pdf(
    ruta_archivo,
    carpeta_temporal
):

    ruta_pdf = os.path.join(
        carpeta_temporal,
        "previsualizacion.pdf"
    )

    convertir_presentacion(
        ruta_archivo,
        ruta_pdf,
        "PDF"
    )

    cantidad = obtener_cantidad_paginas_pdf(
        ruta_pdf
    )

    datos = renderizar_pagina_pdf(
        ruta_pdf,
        0
    )

    return {
        "tipo": "presentacion",
        "modo_presentacion": "pdf",
        "ruta_archivo": ruta_archivo,
        "ruta_pdf": ruta_pdf,
        "rutas_diapositivas": None,
        "pagina_actual": 0,
        "cantidad_paginas": cantidad,
        "datos_imagen": datos,
        "carpeta_temporal": carpeta_temporal
    }


# ================================================================
# PREPARAR PREVISUALIZACIÓN DE PPTX
# ================================================================
def preparar_previsualizacion_presentacion(
    ruta_archivo
):

    carpeta_temporal = tempfile.mkdtemp(
        prefix="convertidor_pptx_"
    )


    # ----------------------------------- #
    #     Preferir PowerPoint directo     #
    # ----------------------------------- #

    if detectar_microsoft_powerpoint():

        try:

            rutas = exportar_diapositivas_powerpoint(
                ruta_archivo,
                carpeta_temporal
            )

            datos = leer_archivo_binario(
                rutas[0]
            )

            return {
                "tipo": "presentacion",
                "modo_presentacion": "imagenes",
                "ruta_archivo": ruta_archivo,
                "ruta_pdf": None,
                "rutas_diapositivas": rutas,
                "pagina_actual": 0,
                "cantidad_paginas": len(rutas),
                "datos_imagen": datos,
                "carpeta_temporal": carpeta_temporal
            }

        except Exception:
            pass


    # ----------------------------------- #
    #       Respaldo mediante PDF         #
    # ----------------------------------- #

    try:

        return preparar_pptx_mediante_pdf(
            ruta_archivo,
            carpeta_temporal
        )

    except Exception:

        shutil.rmtree(
            carpeta_temporal,
            ignore_errors=True
        )

        raise


# ================================================================
# EXTRAER COLOR DE EXCEL
# ================================================================
def extraer_color_excel(
    color
):

    if color is None:
        return None

    try:

        if (
            color.type == "rgb"
            and color.rgb
        ):

            valor = str(
                color.rgb
            )

            # OpenPyXL suele devolver ARGB
            if len(valor) == 8:
                valor = valor[-6:]

            if len(valor) == 6:
                return "#" + valor.upper()

    except Exception:
        pass

    return None


# ================================================================
# CONVERTIR VALOR DE CELDA A TEXTO
# ================================================================
def convertir_valor_celda(
    valor
):

    if valor is None:
        return ""

    if isinstance(
        valor,
        bool
    ):

        return (
            "VERDADERO"
            if valor
            else "FALSO"
        )

    return str(
        valor
    )


# ================================================================
# PREPARAR HOJA DE EXCEL
# ================================================================
def preparar_hoja_excel(
    hoja_valores,
    hoja_formulas
):

    maximo_filas = min(
        max(
            1,
            hoja_valores.max_row,
            hoja_formulas.max_row
        ),
        MAXIMO_FILAS_EXCEL
    )

    maximo_columnas = min(
        max(
            1,
            hoja_valores.max_column,
            hoja_formulas.max_column
        ),
        MAXIMO_COLUMNAS_EXCEL
    )


    # ----------------------------------- #
    #           Celdas                    #
    # ----------------------------------- #

    celdas = []

    for fila in range(
        1,
        maximo_filas + 1
    ):

        for columna in range(
            1,
            maximo_columnas + 1
        ):

            celda_valor = hoja_valores.cell(
                fila,
                columna
            )

            celda_formula = hoja_formulas.cell(
                fila,
                columna
            )

            valor = celda_valor.value


            # Si Excel no guardó un resultado calculado,
            # mostrar al menos la fórmula original
            if (
                valor is None
                and isinstance(
                    celda_formula.value,
                    str
                )
                and celda_formula.value.startswith("=")
            ):

                valor = celda_formula.value


            # No crear datos innecesarios para celdas
            # completamente vacías y sin formato
            if (
                valor is None
                and not celda_formula.has_style
            ):

                continue


            # -------------------------------- #
            #            Fuente                #
            # -------------------------------- #

            fuente = celda_formula.font

            color_fuente = extraer_color_excel(
                fuente.color
            )


            # -------------------------------- #
            #             Fondo                #
            # -------------------------------- #

            color_fondo = None

            relleno = celda_formula.fill

            if relleno.fill_type:

                color_fondo = extraer_color_excel(
                    relleno.fgColor
                )


            # -------------------------------- #
            #          Alineación              #
            # -------------------------------- #

            alineacion = celda_formula.alignment


            celdas.append(
                {
                    "fila": fila - 1,
                    "columna": columna - 1,
                    "texto": convertir_valor_celda(
                        valor
                    ),
                    "negrita": bool(
                        fuente.bold
                    ),
                    "cursiva": bool(
                        fuente.italic
                    ),
                    "subrayado": bool(
                        fuente.underline
                    ),
                    "tamano_fuente": (
                        float(fuente.sz)
                        if fuente.sz
                        else None
                    ),
                    "color_fuente": color_fuente,
                    "color_fondo": color_fondo,
                    "alineacion_horizontal": (
                        alineacion.horizontal
                    ),
                    "alineacion_vertical": (
                        alineacion.vertical
                    )
                }
            )


    # ----------------------------------- #
    #        Columnas y filas             #
    # ----------------------------------- #

    from openpyxl.utils import get_column_letter


    anchos_columnas = {}

    for columna in range(
        1,
        maximo_columnas + 1
    ):

        letra = get_column_letter(
            columna
        )

        dimension = hoja_formulas.column_dimensions[
            letra
        ]

        ancho = dimension.width

        if ancho is not None:

            pixeles = round(
                ancho * 7 + 12
            )

            pixeles = max(
                45,
                min(
                    350,
                    pixeles
                )
            )

            anchos_columnas[
                columna - 1
            ] = pixeles


    altos_filas = {}

    for fila in range(
        1,
        maximo_filas + 1
    ):

        alto = hoja_formulas.row_dimensions[
            fila
        ].height

        if alto is not None:

            pixeles = round(
                alto * 96 / 72
            )

            pixeles = max(
                18,
                min(
                    250,
                    pixeles
                )
            )

            altos_filas[
                fila - 1
            ] = pixeles


    # ----------------------------------- #
    #         Celdas combinadas           #
    # ----------------------------------- #

    combinadas = []

    for rango in hoja_formulas.merged_cells.ranges:

        if (
            rango.min_row > maximo_filas
            or rango.min_col > maximo_columnas
        ):
            continue

        fila = rango.min_row - 1
        columna = rango.min_col - 1

        filas = min(
            rango.max_row,
            maximo_filas
        ) - rango.min_row + 1

        columnas = min(
            rango.max_col,
            maximo_columnas
        ) - rango.min_col + 1

        if filas > 0 and columnas > 0:

            combinadas.append(
                (
                    fila,
                    columna,
                    filas,
                    columnas
                )
            )


    truncada = (
        hoja_formulas.max_row > MAXIMO_FILAS_EXCEL
        or hoja_formulas.max_column > MAXIMO_COLUMNAS_EXCEL
    )


    return {
        "nombre": hoja_formulas.title,
        "filas": maximo_filas,
        "columnas": maximo_columnas,
        "celdas": celdas,
        "anchos_columnas": anchos_columnas,
        "altos_filas": altos_filas,
        "combinadas": combinadas,
        "truncada": truncada
    }


# ================================================================
# PREPARAR PREVISUALIZACIÓN DE XLSX
# ================================================================
def preparar_previsualizacion_hoja_calculo(
    ruta_archivo
):

    try:

        import openpyxl

    except ImportError as error:

        raise RuntimeError(
            "La previsualización de XLSX necesita openpyxl."
        ) from error


    # Libro con resultados calculados
    libro_valores = openpyxl.load_workbook(
        ruta_archivo,
        data_only=True,
        read_only=False,
        keep_links=False
    )

    # Libro con fórmulas y estilos
    libro_formulas = openpyxl.load_workbook(
        ruta_archivo,
        data_only=False,
        read_only=False,
        keep_links=False
    )


    try:

        hojas = []


        # ----------------------------------- #
        #      Mostrar hojas visibles         #
        # ----------------------------------- #

        for nombre in libro_formulas.sheetnames:

            hoja_formulas = libro_formulas[
                nombre
            ]

            if hoja_formulas.sheet_state != "visible":
                continue

            hoja_valores = libro_valores[
                nombre
            ]

            hojas.append(
                preparar_hoja_excel(
                    hoja_valores,
                    hoja_formulas
                )
            )


        # Si excepcionalmente todas están ocultas,
        # mostrar al menos la hoja activa
        if not hojas:

            hoja_formulas = libro_formulas.active
            hoja_valores = libro_valores[
                hoja_formulas.title
            ]

            hojas.append(
                preparar_hoja_excel(
                    hoja_valores,
                    hoja_formulas
                )
            )


        return {
            "tipo": "hoja_calculo",
            "ruta_archivo": ruta_archivo,
            "hojas": hojas,
            "carpeta_temporal": None
        }


    finally:

        libro_valores.close()
        libro_formulas.close()


# ================================================================
# PREPARAR PREVISUALIZACIÓN
# ================================================================
def preparar_previsualizacion(
    ruta_archivo
):

    tipo = obtener_tipo_previsualizacion(
        ruta_archivo
    )


    if tipo == "imagen":

        return preparar_previsualizacion_imagen(
            ruta_archivo
        )


    if tipo == "documento":

        return preparar_previsualizacion_documento(
            ruta_archivo
        )


    if tipo == "presentacion":

        return preparar_previsualizacion_presentacion(
            ruta_archivo
        )


    if tipo == "hoja_calculo":

        return preparar_previsualizacion_hoja_calculo(
            ruta_archivo
        )


    return {
        "tipo": "archivo",
        "ruta_archivo": ruta_archivo,
        "carpeta_temporal": None
    }


# ================================================================
# OBTENER IMAGEN DE PÁGINA
# ================================================================
def obtener_imagen_pagina_previsualizacion(
    resultado,
    indice_pagina
):

    tipo = resultado[
        "tipo"
    ]


    # ----------------------------------- #
    #          PDF o DOCX                 #
    # ----------------------------------- #

    if tipo == "documento":

        return renderizar_pagina_pdf(
            resultado["ruta_pdf"],
            indice_pagina
        )


    # ----------------------------------- #
    #              PPTX                   #
    # ----------------------------------- #

    if tipo == "presentacion":

        if resultado[
            "modo_presentacion"
        ] == "imagenes":

            ruta = resultado[
                "rutas_diapositivas"
            ][
                indice_pagina
            ]

            return leer_archivo_binario(
                ruta
            )

        return renderizar_pagina_pdf(
            resultado["ruta_pdf"],
            indice_pagina
        )


    raise RuntimeError(
        "Este tipo de previsualización no utiliza páginas."
    )


# ================================================================
# LIMPIAR PREVISUALIZACIÓN
# ================================================================
def limpiar_previsualizacion(
    resultado
):

    if not resultado:
        return

    carpeta_temporal = resultado.get(
        "carpeta_temporal"
    )

    if carpeta_temporal:

        shutil.rmtree(
            carpeta_temporal,
            ignore_errors=True
        )
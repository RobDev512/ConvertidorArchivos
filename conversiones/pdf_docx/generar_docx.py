import json
import re
import statistics
import zipfile
from copy import deepcopy
from pathlib import Path

from lxml import etree

from docx import Document
from docx.enum.table import (
    WD_CELL_VERTICAL_ALIGNMENT,
    WD_ROW_HEIGHT_RULE,
    WD_TABLE_ALIGNMENT
)
from docx.enum.text import (
    WD_ALIGN_PARAGRAPH,
    WD_LINE_SPACING
)
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor


# ================================================================
# RUTAS Y CONFIGURACION
# ================================================================

CARPETA_BASE = Path(__file__).resolve().parent
CARPETA_SALIDA = CARPETA_BASE / "salida"

RUTA_DISTRIBUCION = CARPETA_SALIDA / "distribucion_pdf.json"
RUTA_EFECTOS_REFINADOS = CARPETA_SALIDA / "analisis_efectos_pdf_refinado.json"
RUTA_ESPACIADOS_REFINADOS = CARPETA_SALIDA / "espaciados_pdf_refinados.json"
RUTA_DOCX = CARPETA_SALIDA / "resultado_convertido.docx"

NS_VML = "urn:schemas-microsoft-com:vml"
NS_OFFICE = "urn:schemas-microsoft-com:office:office"
NS_REL = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"

FUENTES_EQUIVALENTES = {
    "ArialMT": "Arial",
    "Arial-BoldMT": "Arial",
    "Arial-ItalicMT": "Arial",
    "Arial-BoldItalicMT": "Arial",
    "Aptos": "Aptos",
    "Aptos,Bold": "Aptos",
    "Aptos,Italic": "Aptos",
    "Aptos,BoldItalic": "Aptos",
    "BellMT": "Bell MT",
    "BellMTBold": "Bell MT",
    "BellMTItalic": "Bell MT",
    "Cambria": "Cambria",
    "Cambria-Bold": "Cambria",
    "Cambria-Italic": "Cambria",
    "Cambria-BoldItalic": "Cambria",
    "SymbolMT": "Symbol"
}


# ================================================================
# CARGA Y UTILIDADES
# ================================================================

def cargar_json(ruta):
    if not ruta.is_file():
        raise FileNotFoundError(
            f"No existe el archivo requerido: {ruta}"
        )

    with open(ruta, "r", encoding="utf-8") as archivo:
        return json.load(archivo)


def cargar_json_opcional(ruta):
    if not ruta.is_file():
        return {}

    with open(ruta, "r", encoding="utf-8") as archivo:
        return json.load(archivo)


def numero(valor, predeterminado=0.0):
    try:
        return float(valor)
    except (TypeError, ValueError):
        return float(predeterminado)


def mediana(valores, predeterminado=0.0):
    validos = []

    for valor in valores:
        try:
            validos.append(float(valor))
        except (TypeError, ValueError):
            pass

    if not validos:
        return float(predeterminado)

    return float(statistics.median(validos))


def limitar(valor, minimo, maximo):
    return max(
        minimo,
        min(
            maximo,
            valor
        )
    )


# ================================================================
# TEXTO Y RECURSOS
# ================================================================

def normalizar_texto_clave(texto):
    return re.sub(
        r"\s+",
        " ",
        str(texto or "")
    ).strip().casefold()


def limpiar_texto_indice(texto):
    texto = re.sub(
        r"\s+",
        " ",
        str(texto or "")
    ).strip()

    texto = re.sub(
        r"\s*_+\s*\d*\s*$",
        "",
        texto
    ).rstrip()

    return texto


def obtener_pagina_indice(elemento):
    pagina = elemento.get(
        "pagina_destino"
    )

    if pagina is not None:
        return str(pagina)

    texto = str(
        elemento.get(
            "texto",
            ""
        )
    )

    coincidencia = re.search(
        r"_+\s*(\d+)\s*$",
        texto
    )

    if coincidencia:
        return coincidencia.group(1)

    return ""


def resolver_ruta_recurso(ruta_guardada):
    if not ruta_guardada:
        return None

    texto_ruta = str(
        ruta_guardada
    )

    ruta = Path(
        texto_ruta
    )

    if ruta.is_file():
        return ruta

    partes = [
        parte
        for parte in re.split(
            r"[\\/]+",
            texto_ruta
        )
        if parte
    ]

    for carpeta in (
        "recursos_fondo",
        "recursos_paginas"
    ):
        if carpeta not in partes:
            continue

        indice = partes.index(
            carpeta
        )

        candidato = CARPETA_SALIDA.joinpath(
            *partes[indice:]
        )

        if candidato.is_file():
            return candidato

    return None


# ================================================================
# FUENTES Y COLORES
# ================================================================

def normalizar_fuente(nombre):
    nombre = str(
        nombre or ""
    ).strip()

    if nombre in FUENTES_EQUIVALENTES:
        return FUENTES_EQUIVALENTES[
            nombre
        ]

    if "," in nombre:
        base = nombre.split(
            ",",
            1
        )[0].strip()

        if base:
            return FUENTES_EQUIVALENTES.get(
                base,
                base
            )

    for sufijo in (
        "-BoldItalicMT",
        "-BoldMT",
        "-ItalicMT",
        "MT",
        "-BoldItalic",
        "-Bold",
        "-Italic"
    ):
        if nombre.endswith(
            sufijo
        ):
            base = nombre[
                :-len(sufijo)
            ]

            if base:
                return FUENTES_EQUIVALENTES.get(
                    base,
                    base
                )

    return nombre or "Aptos"


def color_pdf_a_tupla(color):
    if isinstance(
        color,
        (int, float)
    ):
        valor = int(
            round(
                limitar(
                    float(color),
                    0.0,
                    1.0
                )
                * 255.0
            )
        )

        return (
            valor,
            valor,
            valor
        )

    if (
        isinstance(
            color,
            (list, tuple)
        )
        and color
    ):
        valores = list(
            color
        )

        if len(valores) == 1:
            valores *= 3

        if len(valores) >= 3:
            return tuple(
                int(
                    round(
                        limitar(
                            numero(valor),
                            0.0,
                            1.0
                        )
                        * 255.0
                    )
                )
                for valor in valores[:3]
            )

    return (
        0,
        0,
        0
    )


def convertir_color_pdf(color):
    return RGBColor(
        *color_pdf_a_tupla(
            color
        )
    )


def rgb_lista_a_tupla(
    rgb,
    predeterminado=(0, 0, 0)
):
    if (
        not isinstance(
            rgb,
            (list, tuple)
        )
        or len(rgb) < 3
    ):
        return predeterminado

    return tuple(
        int(
            round(
                limitar(
                    numero(valor),
                    0.0,
                    255.0
                )
            )
        )
        for valor in rgb[:3]
    )


def rgb_a_hex(rgb):
    r, g, b = rgb_lista_a_tupla(
        rgb
    )

    return (
        f"{r:02X}"
        f"{g:02X}"
        f"{b:02X}"
    )


def convertir_alineacion(alineacion):
    mapa = {
        "izquierda": WD_ALIGN_PARAGRAPH.LEFT,
        "centro": WD_ALIGN_PARAGRAPH.CENTER,
        "derecha": WD_ALIGN_PARAGRAPH.RIGHT,
        "justificada": WD_ALIGN_PARAGRAPH.JUSTIFY
    }

    return mapa.get(
        str(
            alineacion or ""
        ).lower(),
        WD_ALIGN_PARAGRAPH.LEFT
    )


# ================================================================
# EFECTOS DETECTADOS
# ================================================================

def construir_mapa_efectos(
    datos_efectos
):
    por_id = {}
    por_texto = {}

    for candidato in datos_efectos.get(
        "candidatos",
        []
    ):
        pagina = int(
            numero(
                candidato.get(
                    "pagina"
                ),
                0
            )
        )

        identificador = str(
            candidato.get(
                "id"
            )
            or ""
        )

        texto = normalizar_texto_clave(
            candidato.get(
                "texto"
            )
        )

        if pagina <= 0:
            continue

        if identificador:
            por_id[
                (
                    pagina,
                    identificador
                )
            ] = candidato

        if texto:
            por_texto.setdefault(
                (
                    pagina,
                    texto
                ),
                []
            ).append(
                candidato
            )

    return {
        "por_id": por_id,
        "por_texto": por_texto
    }


def obtener_efecto(
    mapa_efectos,
    numero_pagina,
    identificador=None,
    texto=None
):
    pagina = int(
        numero(
            numero_pagina,
            0
        )
    )

    if identificador:
        efecto = mapa_efectos[
            "por_id"
        ].get(
            (
                pagina,
                str(
                    identificador
                )
            )
        )

        if efecto:
            return efecto

    clave_texto = normalizar_texto_clave(
        texto
    )

    if clave_texto:
        candidatos = mapa_efectos[
            "por_texto"
        ].get(
            (
                pagina,
                clave_texto
            ),
            []
        )

        if len(
            candidatos
        ) == 1:
            return candidatos[0]

    return None


def nombres_imagenes_reflejo(
    mapa_efectos,
    numero_pagina=None
):
    nombres = set()

    for (
        pagina,
        _
    ), efecto in mapa_efectos[
        "por_id"
    ].items():
        if (
            numero_pagina is not None
            and pagina
            != int(
                numero_pagina
            )
        ):
            continue

        reflejo = efecto.get(
            "reflejo_final",
            {}
        )

        imagen = (
            reflejo.get(
                "imagen"
            )
            or {}
        )

        nombre = imagen.get(
            "nombre"
        )

        if nombre:
            nombres.add(
                str(
                    nombre
                )
            )

    return nombres


# ================================================================
# ESPACIADOS REFINADOS DETECTADOS DESDE EL PDF
# ================================================================

def construir_mapa_espaciados(
    datos_espaciados
):
    por_indice = {}
    por_texto = {}

    for pagina in datos_espaciados.get(
        "paginas",
        []
    ):
        numero_pagina = int(
            numero(
                pagina.get(
                    "numero"
                ),
                0
            )
        )

        if numero_pagina <= 0:
            continue

        for elemento in pagina.get(
            "elementos",
            []
        ):
            indice = int(
                numero(
                    elemento.get(
                        "indice"
                    ),
                    -1
                )
            )

            tipo = str(
                elemento.get(
                    "tipo"
                )
                or ""
            )

            texto = normalizar_texto_clave(
                elemento.get(
                    "texto"
                )
            )

            if indice >= 0:
                por_indice[
                    (
                        numero_pagina,
                        indice
                    )
                ] = elemento

            if texto:
                por_texto.setdefault(
                    (
                        numero_pagina,
                        tipo,
                        texto
                    ),
                    []
                ).append(
                    elemento
                )

    return {
        "por_indice": por_indice,
        "por_texto": por_texto
    }


def obtener_espaciado_refinado(
    mapa_espaciados,
    numero_pagina,
    indice_elemento=None,
    tipo=None,
    texto=None
):
    if not mapa_espaciados:
        return None

    pagina = int(
        numero(
            numero_pagina,
            0
        )
    )

    if (
        pagina <= 0
        or not isinstance(
            mapa_espaciados,
            dict
        )
    ):
        return None

    if indice_elemento is not None:
        try:
            indice = int(
                indice_elemento
            )
        except (
            TypeError,
            ValueError
        ):
            indice = None

        if indice is not None:
            encontrado = mapa_espaciados.get(
                "por_indice",
                {}
            ).get(
                (
                    pagina,
                    indice
                )
            )

            if encontrado is not None:
                return encontrado

    clave_texto = normalizar_texto_clave(
        texto
    )

    if clave_texto:
        candidatos = mapa_espaciados.get(
            "por_texto",
            {}
        ).get(
            (
                pagina,
                str(
                    tipo
                    or ""
                ),
                clave_texto
            ),
            []
        )

        if len(
            candidatos
        ) == 1:
            return candidatos[
                0
            ]

    return None


def resolver_espaciado_word(
    espaciado_refinado,
    elemento,
    tipo_pagina,
    margenes_word,
    es_primero
):
    if isinstance(
        espaciado_refinado,
        dict
    ):
        antes = max(
            0.0,
            numero(
                espaciado_refinado.get(
                    "espacio_word_antes"
                )
            )
        )

        despues = max(
            0.0,
            numero(
                espaciado_refinado.get(
                    "espacio_word_despues"
                )
            )
        )

        return (
            antes,
            despues,
            True
        )

    antes, despues = inferir_espaciado_parrafo(
        elemento,
        tipo_pagina,
        margenes_word,
        es_primero
    )

    return (
        antes,
        despues,
        False
    )


# ================================================================
# MARGENES DEL DOCUMENTO
# ================================================================

def inferir_margenes_word(
    distribucion
):
    paginas = distribucion.get(
        "paginas",
        []
    )

    declarados = distribucion.get(
        "margenes_documento",
        {}
    )

    ancho_pagina = numero(
        paginas[0].get(
            "ancho"
        )
        if paginas
        else 612.0,
        612.0
    )

    candidatos_izq = []
    candidatos_der = []

    for pagina in paginas:
        if (
            pagina.get(
                "tipo"
            )
            == "portada"
        ):
            continue

        for elemento in pagina.get(
            "elementos",
            []
        ):
            caja = elemento.get(
                "caja",
                {}
            )

            x0 = numero(
                caja.get(
                    "x0"
                ),
                -1.0
            )

            x1 = numero(
                caja.get(
                    "x1"
                ),
                -1.0
            )

            if (
                24.0
                <= x0
                <= ancho_pagina / 2.0
            ):
                candidatos_izq.append(
                    x0
                )

            if x1 > ancho_pagina / 2.0:
                margen_derecho = (
                    ancho_pagina
                    - x1
                )

                if (
                    24.0
                    <= margen_derecho
                    <= ancho_pagina / 2.0
                ):
                    candidatos_der.append(
                        margen_derecho
                    )

    margen_izquierdo = numero(
        declarados.get(
            "izquierdo"
        ),
        72.0
    )

    margen_derecho = numero(
        declarados.get(
            "derecho"
        ),
        72.0
    )

    if candidatos_izq:
        minimo_izq = min(
            candidatos_izq
        )

        if (
            36.0
            <= minimo_izq
            <= 144.0
        ):
            margen_izquierdo = (
                minimo_izq
            )

    if candidatos_der:
        minimo_der = min(
            candidatos_der
        )

        if (
            36.0
            <= minimo_der
            <= 144.0
        ):
            margen_derecho = (
                minimo_der
            )

    medidas_comunes = (
        18.0,
        36.0,
        54.0,
        72.0,
        90.0,
        108.0,
        144.0
    )

    def ajustar_medida(valor):
        cercana = min(
            medidas_comunes,
            key=lambda medida: abs(
                medida
                - valor
            )
        )

        if abs(
            cercana
            - valor
        ) <= 2.0:
            return cercana

        return round(
            valor
            * 4.0
        ) / 4.0

    return {
        "izquierdo": ajustar_medida(
            margen_izquierdo
        ),
        "derecho": ajustar_medida(
            margen_derecho
        ),
        "superior": numero(
            declarados.get(
                "superior"
            ),
            72.0
        ),
        "inferior": numero(
            declarados.get(
                "inferior"
            ),
            72.0
        )
    }


# ================================================================
# FORMATO DE TEXTO
# ================================================================

def desactivar_ajuste_cuadricula(
    parrafo
):
    p_pr = (
        parrafo
        ._p
        .get_or_add_pPr()
    )

    snap = p_pr.find(
        qn(
            "w:snapToGrid"
        )
    )

    if snap is None:
        snap = OxmlElement(
            "w:snapToGrid"
        )

        p_pr.append(
            snap
        )

    snap.set(
        qn(
            "w:val"
        ),
        "0"
    )


def aplicar_formato_run(
    run,
    datos,
    estilo_predeterminado=None
):
    estilo_predeterminado = (
        estilo_predeterminado
        or {}
    )

    fuente = (
        datos.get(
            "fuente"
        )
        or estilo_predeterminado.get(
            "fuente"
        )
    )

    tamano = datos.get(
        "tamano"
    )

    color = datos.get(
        "color"
    )

    negrita = datos.get(
        "negrita"
    )

    cursiva = datos.get(
        "cursiva"
    )

    if tamano is None:
        tamano = estilo_predeterminado.get(
            "tamano"
        )

    if color is None:
        color = estilo_predeterminado.get(
            "color"
        )

    if negrita is None:
        negrita = estilo_predeterminado.get(
            "negrita",
            False
        )

    if cursiva is None:
        cursiva = estilo_predeterminado.get(
            "cursiva",
            False
        )

    run.font.name = normalizar_fuente(
        fuente
    )

    if tamano is not None:
        run.font.size = Pt(
            max(
                1.0,
                numero(
                    tamano,
                    11.0
                )
            )
        )

    run.bold = bool(
        negrita
    )

    run.italic = bool(
        cursiva
    )

    run.font.color.rgb = convertir_color_pdf(
        color
    )


def tamano_decorativo(elemento):
    estilo = elemento.get(
        "estilo",
        {}
    )

    tamano = numero(
        estilo.get(
            "tamano"
        ),
        11.0
    )

    tipo = str(
        elemento.get(
            "tipo"
        )
        or ""
    )

    caja = elemento.get(
        "caja",
        {}
    )

    alto = (
        numero(
            caja.get(
                "bottom"
            )
        )
        - numero(
            caja.get(
                "top"
            )
        )
    )

    if (
        tipo in {
            "titulo_portada_principal",
            "titulo_portada",
            "encabezado"
        }
        and 8.0 <= alto <= 120.0
    ):
        return alto

    return tamano


# ================================================================
# EFECTOS OOXML W14
# ================================================================

def limpiar_efectos_run(run):
    r_pr = (
        run
        ._r
        .get_or_add_rPr()
    )

    for nombre in (
        "w14:textFill",
        "w14:reflection"
    ):
        elemento = r_pr.find(
            qn(
                nombre
            )
        )

        if elemento is not None:
            r_pr.remove(
                elemento
            )


def aplicar_relleno_solido_w14(
    run,
    rgb
):
    r_pr = (
        run
        ._r
        .get_or_add_rPr()
    )

    text_fill = OxmlElement(
        "w14:textFill"
    )

    solid_fill = OxmlElement(
        "w14:solidFill"
    )

    srgb = OxmlElement(
        "w14:srgbClr"
    )

    srgb.set(
        qn(
            "w14:val"
        ),
        rgb_a_hex(
            rgb
        )
    )

    solid_fill.append(
        srgb
    )

    text_fill.append(
        solid_fill
    )

    r_pr.append(
        text_fill
    )


def aplicar_degradado_w14(
    run,
    relleno
):
    stops = relleno.get(
        "stops",
        []
    )

    if not stops:
        return False

    r_pr = (
        run
        ._r
        .get_or_add_rPr()
    )

    text_fill = OxmlElement(
        "w14:textFill"
    )

    grad_fill = OxmlElement(
        "w14:gradFill"
    )

    gs_lst = OxmlElement(
        "w14:gsLst"
    )

    cantidad_stops = 0

    for stop in stops:
        rgb = stop.get(
            "rgb"
        )

        if not rgb:
            continue

        gs = OxmlElement(
            "w14:gs"
        )

        gs.set(
            qn(
                "w14:pos"
            ),
            str(
                int(
                    round(
                        limitar(
                            numero(
                                stop.get(
                                    "pos"
                                )
                            ),
                            0.0,
                            100000.0
                        )
                    )
                )
            )
        )

        srgb = OxmlElement(
            "w14:srgbClr"
        )

        srgb.set(
            qn(
                "w14:val"
            ),
            rgb_a_hex(
                rgb
            )
        )

        gs.append(
            srgb
        )

        gs_lst.append(
            gs
        )

        cantidad_stops += 1

    if cantidad_stops == 0:
        return False

    grad_fill.append(
        gs_lst
    )

    lin = OxmlElement(
        "w14:lin"
    )

    direccion = str(
        relleno.get(
            "direccion"
        )
        or "vertical"
    ).lower()

    if direccion == "horizontal":
        angulo = 0

    elif direccion == "diagonal":
        angulo = 2700000

    else:
        angulo = 5400000

    lin.set(
        qn(
            "w14:ang"
        ),
        str(
            angulo
        )
    )

    lin.set(
        qn(
            "w14:scaled"
        ),
        "0"
    )

    grad_fill.append(
        lin
    )

    text_fill.append(
        grad_fill
    )

    r_pr.append(
        text_fill
    )

    return True


def aplicar_reflejo_w14(
    run,
    reflejo
):
    if not reflejo.get(
        "presente"
    ):
        return False

    imagen = (
        reflejo.get(
            "imagen"
        )
        or {}
    )

    alpha_max = numero(
        imagen.get(
            "alpha_max"
        ),
        128.0
    )

    inicio_alpha = int(
        round(
            limitar(
                (
                    alpha_max
                    / 255.0
                    * 100000.0
                ),
                25000.0,
                75000.0
            )
        )
    )

    r_pr = (
        run
        ._r
        .get_or_add_rPr()
    )

    nodo = OxmlElement(
        "w14:reflection"
    )

    atributos = {
        "blurRad": "6350",
        "stA": str(
            inicio_alpha
        ),
        "stPos": "0",
        "endA": "300",
        "endPos": "40000",
        "dist": "0",
        "dir": "5400000",
        "fadeDir": "5400000",
        "sx": "100000",
        "sy": "-95000",
        "kx": "0",
        "ky": "0",
        "algn": "bl"
    }

    for clave, valor in atributos.items():
        nodo.set(
            qn(
                f"w14:{clave}"
            ),
            valor
        )

    r_pr.append(
        nodo
    )

    return True


def aplicar_efecto_run(
    run,
    efecto,
    estadisticas
):
    if not efecto:
        return False

    limpiar_efectos_run(
        run
    )

    color_final = efecto.get(
        "color_final",
        {}
    )

    rgb = color_final.get(
        "rgb"
    )

    if rgb:
        run.font.color.rgb = RGBColor(
            *rgb_lista_a_tupla(
                rgb
            )
        )

    relleno = efecto.get(
        "relleno_final",
        {}
    )

    aplico_text_fill = False
    aplico_reflejo = False

    if relleno.get(
        "tipo"
    ) == "degradado":
        aplico_text_fill = aplicar_degradado_w14(
            run,
            relleno
        )

    elif (
        rgb
        and color_final.get(
            "origen"
        )
        == "analisis_visual"
    ):
        aplicar_relleno_solido_w14(
            run,
            rgb
        )

        aplico_text_fill = True

    aplico_reflejo = aplicar_reflejo_w14(
        run,
        efecto.get(
            "reflejo_final",
            {}
        )
    )

    if (
        rgb
        or aplico_text_fill
        or aplico_reflejo
    ):
        estadisticas[
            "efectos_aplicados"
        ] += 1

    if aplico_text_fill:
        estadisticas[
            "text_fill_aplicados"
        ] += 1

    if aplico_reflejo:
        estadisticas[
            "reflejos_aplicados"
        ] += 1

    return bool(
        rgb
        or aplico_text_fill
        or aplico_reflejo
    )


def parece_continuacion_url(
    texto_acumulado,
    texto_actual
):
    acumulado = str(
        texto_acumulado
        or ""
    )

    actual = str(
        texto_actual
        or ""
    ).strip()

    if not actual:
        return False

    indice_http = max(
        acumulado.rfind(
            "http://"
        ),
        acumulado.rfind(
            "https://"
        )
    )

    if indice_http < 0:
        return False

    sufijo = acumulado[
        indice_http:
    ]

    if any(
        caracter.isspace()
        for caracter in sufijo
    ):
        return False

    permitidos = set(
        (
            "abcdefghijklmnopqrstuvwxyz"
            "ABCDEFGHIJKLMNOPQRSTUVWXYZ"
            "0123456789"
            "%/?=&._-:#~+@,;()[]!$'"
        )
    )

    return all(
        caracter in permitidos
        for caracter in actual
    )


def separador_entre_lineas(
    texto_acumulado,
    texto_anterior,
    texto_actual
):
    anterior = str(
        texto_anterior
        or ""
    ).strip()

    actual = str(
        texto_actual
        or ""
    ).strip()

    if (
        not anterior
        or not actual
    ):
        return ""

    if parece_continuacion_url(
        texto_acumulado,
        actual
    ):
        return ""

    if (
        anterior.endswith(
            "-"
        )
        and actual[:1].islower()
    ):
        return ""

    return " "


# ================================================================
# CONSTRUCCION DE TEXTO
# ================================================================

def agregar_lineas_formateadas(
    parrafo,
    lineas,
    estilo_predeterminado=None
):
    estilo_predeterminado = (
        estilo_predeterminado
        or {}
    )

    texto_anterior = ""
    texto_acumulado = ""
    hay_contenido = False

    for linea in lineas:
        palabras = linea.get(
            "palabras",
            []
        )

        texto_linea = str(
            linea.get(
                "texto",
                ""
            )
        ).strip()

        if hay_contenido:
            separador = separador_entre_lineas(
                texto_acumulado,
                texto_anterior,
                texto_linea
            )

            parrafo.add_run(
                separador
            )

            texto_acumulado += separador

        if palabras:
            primera_palabra = True

            for palabra in palabras:
                texto = str(
                    palabra.get(
                        "texto",
                        ""
                    )
                )

                if not texto:
                    continue

                if not primera_palabra:
                    parrafo.add_run(
                        " "
                    )

                    texto_acumulado += " "

                run = parrafo.add_run(
                    texto
                )

                aplicar_formato_run(
                    run,
                    palabra,
                    estilo_predeterminado
                )

                texto_acumulado += texto
                primera_palabra = False

        elif texto_linea:
            run = parrafo.add_run(
                texto_linea
            )

            aplicar_formato_run(
                run,
                {
                    "fuente": linea.get(
                        "fuente_predominante"
                    ),
                    "tamano": linea.get(
                        "tamano_promedio"
                    ),
                    "negrita": linea.get(
                        "negrita_predominante"
                    ),
                    "cursiva": linea.get(
                        "cursiva_predominante"
                    ),
                    "color": linea.get(
                        "color_predominante"
                    )
                },
                estilo_predeterminado
            )

            texto_acumulado += texto_linea

        if (
            palabras
            or texto_linea
        ):
            hay_contenido = True
            texto_anterior = texto_linea


# ================================================================
# IMAGENES Y FONDO
# ================================================================

def agregar_imagen_flotante(
    parrafo,
    ruta_imagen,
    x,
    y,
    ancho,
    alto,
    detras_texto=True
):
    ruta_imagen = (
        resolver_ruta_recurso(
            ruta_imagen
        )
        or Path(
            str(
                ruta_imagen
            )
        )
    )

    if not ruta_imagen.is_file():
        return False

    run = parrafo.add_run()

    run.add_picture(
        str(
            ruta_imagen
        ),
        width=Pt(
            max(
                0.1,
                numero(
                    ancho
                )
            )
        ),
        height=Pt(
            max(
                0.1,
                numero(
                    alto
                )
            )
        )
    )

    drawing = run._r.find(
        qn(
            "w:drawing"
        )
    )

    inline = (
        drawing.find(
            qn(
                "wp:inline"
            )
        )
        if drawing is not None
        else None
    )

    if inline is None:
        return False

    extent = inline.find(
        qn(
            "wp:extent"
        )
    )

    doc_pr = inline.find(
        qn(
            "wp:docPr"
        )
    )

    c_nv = inline.find(
        qn(
            "wp:cNvGraphicFramePr"
        )
    )

    graphic = inline.find(
        qn(
            "a:graphic"
        )
    )

    if any(
        elemento is None
        for elemento in (
            extent,
            doc_pr,
            c_nv,
            graphic
        )
    ):
        return False

    anchor = OxmlElement(
        "wp:anchor"
    )

    atributos = {
        "distT": "0",
        "distB": "0",
        "distL": "0",
        "distR": "0",
        "simplePos": "0",
        "relativeHeight": (
            "0"
            if detras_texto
            else "251658240"
        ),
        "behindDoc": (
            "1"
            if detras_texto
            else "0"
        ),
        "locked": "0",
        "layoutInCell": "0",
        "allowOverlap": "1"
    }

    for clave, valor in atributos.items():
        anchor.set(
            clave,
            valor
        )

    simple_pos = OxmlElement(
        "wp:simplePos"
    )

    simple_pos.set(
        "x",
        "0"
    )

    simple_pos.set(
        "y",
        "0"
    )

    anchor.append(
        simple_pos
    )

    position_h = OxmlElement(
        "wp:positionH"
    )

    position_h.set(
        "relativeFrom",
        "page"
    )

    offset_h = OxmlElement(
        "wp:posOffset"
    )

    offset_h.text = str(
        Pt(
            numero(
                x
            )
        ).emu
    )

    position_h.append(
        offset_h
    )

    anchor.append(
        position_h
    )

    position_v = OxmlElement(
        "wp:positionV"
    )

    position_v.set(
        "relativeFrom",
        "page"
    )

    offset_v = OxmlElement(
        "wp:posOffset"
    )

    offset_v.text = str(
        Pt(
            numero(
                y
            )
        ).emu
    )

    position_v.append(
        offset_v
    )

    anchor.append(
        position_v
    )

    anchor.append(
        deepcopy(
            extent
        )
    )

    effect_extent = OxmlElement(
        "wp:effectExtent"
    )

    for lado in (
        "l",
        "t",
        "r",
        "b"
    ):
        effect_extent.set(
            lado,
            "0"
        )

    anchor.append(
        effect_extent
    )

    anchor.append(
        OxmlElement(
            "wp:wrapNone"
        )
    )

    anchor.append(
        deepcopy(
            doc_pr
        )
    )

    anchor.append(
        deepcopy(
            c_nv
        )
    )

    anchor.append(
        deepcopy(
            graphic
        )
    )

    drawing.remove(
        inline
    )

    drawing.append(
        anchor
    )

    return True


def agregar_fondo_mosaico_vml(
    story,
    ruta_textura,
    ancho_pagina,
    alto_pagina
):
    ruta_textura = resolver_ruta_recurso(
        ruta_textura
    )

    if (
        ruta_textura is None
        or not ruta_textura.is_file()
    ):
        return False

    rid, _ = story.part.get_or_add_image(
        str(
            ruta_textura
        )
    )

    parrafo = story.paragraphs[
        0
    ]

    run = parrafo.add_run()

    pict = OxmlElement(
        "w:pict"
    )

    rect = etree.Element(
        f"{{{NS_VML}}}rect",
        nsmap={
            "v": NS_VML,
            "o": NS_OFFICE,
            "r": NS_REL
        }
    )

    rect.set(
        "id",
        "_x0000_s1025"
    )

    rect.set(
        "style",
        (
            "position:absolute;"
            "margin-left:0pt;"
            "margin-top:0pt;"
            f"width:{ancho_pagina:.3f}pt;"
            f"height:{alto_pagina:.3f}pt;"
            "z-index:-251658251;"
            "mso-position-horizontal-relative:page;"
            "mso-position-vertical-relative:page"
        )
    )

    rect.set(
        "stroked",
        "f"
    )

    fill = etree.SubElement(
        rect,
        f"{{{NS_VML}}}fill"
    )

    fill.set(
        f"{{{NS_REL}}}id",
        rid
    )

    fill.set(
        "type",
        "tile"
    )

    fill.set(
        "recolor",
        "f"
    )

    fill.set(
        "rotate",
        "t"
    )

    stroke = etree.SubElement(
        rect,
        f"{{{NS_VML}}}stroke"
    )

    stroke.set(
        "on",
        "f"
    )

    pict.append(
        rect
    )

    run._r.append(
        pict
    )

    return True


def obtener_textura_fondo(
    distribucion
):
    fondo = distribucion.get(
        "fondo_documento",
        {}
    )

    if not fondo.get(
        "detectado"
    ):
        return None

    for archivo in fondo.get(
        "archivos_extraidos",
        []
    ):
        ruta = resolver_ruta_recurso(
            archivo
        )

        if ruta:
            return ruta

    return None


# ================================================================
# CONFIGURACION DE PAGINA
# ================================================================

def configurar_borde_pagina(
    section,
    datos_borde
):
    sect_pr = section._sectPr

    anterior = sect_pr.find(
        qn(
            "w:pgBorders"
        )
    )

    if anterior is not None:
        sect_pr.remove(
            anterior
        )

    pg_borders = OxmlElement(
        "w:pgBorders"
    )

    pg_borders.set(
        qn(
            "w:offsetFrom"
        ),
        "page"
    )

    pg_borders.set(
        qn(
            "w:display"
        ),
        "allPages"
    )

    nombres = {
        "top": "superior",
        "left": "izquierdo",
        "bottom": "inferior",
        "right": "derecho"
    }

    for lado, nombre in nombres.items():
        datos = datos_borde.get(
            nombre,
            {}
        )

        if not datos.get(
            "presente",
            False
        ):
            continue

        grosor = numero(
            datos.get(
                "grosor"
            ),
            0.4
        )

        sz = int(
            round(
                grosor
                * 8.0
            )
        )

        if (
            0.25
            <= grosor
            <= 0.65
        ):
            sz = max(
                4,
                sz
            )

        else:
            sz = max(
                2,
                sz
            )

        borde = OxmlElement(
            f"w:{lado}"
        )

        borde.set(
            qn(
                "w:val"
            ),
            "dotted"
        )

        borde.set(
            qn(
                "w:sz"
            ),
            str(
                sz
            )
        )

        borde.set(
            qn(
                "w:space"
            ),
            str(
                max(
                    0,
                    int(
                        round(
                            numero(
                                datos.get(
                                    "separacion_borde"
                                ),
                                10.0
                            )
                        )
                    )
                )
            )
        )

        borde.set(
            qn(
                "w:color"
            ),
            "auto"
        )

        pg_borders.append(
            borde
        )

    sect_pr.append(
        pg_borders
    )


def configurar_numero_pagina_inicio(
    section,
    inicio=0
):
    sect_pr = section._sectPr

    pg_num_type = sect_pr.find(
        qn(
            "w:pgNumType"
        )
    )

    if pg_num_type is None:
        pg_num_type = OxmlElement(
            "w:pgNumType"
        )

        sect_pr.append(
            pg_num_type
        )

    pg_num_type.set(
        qn(
            "w:start"
        ),
        str(
            int(
                inicio
            )
        )
    )


def agregar_campo_page(
    parrafo,
    negrita=True
):
    fld_simple = OxmlElement(
        "w:fldSimple"
    )

    fld_simple.set(
        qn(
            "w:instr"
        ),
        " PAGE "
    )

    run_xml = OxmlElement(
        "w:r"
    )

    r_pr = OxmlElement(
        "w:rPr"
    )

    r_fonts = OxmlElement(
        "w:rFonts"
    )

    r_fonts.set(
        qn(
            "w:ascii"
        ),
        "Aptos"
    )

    r_fonts.set(
        qn(
            "w:hAnsi"
        ),
        "Aptos"
    )

    r_pr.append(
        r_fonts
    )

    sz = OxmlElement(
        "w:sz"
    )

    sz.set(
        qn(
            "w:val"
        ),
        "24"
    )

    r_pr.append(
        sz
    )

    if negrita:
        r_pr.append(
            OxmlElement(
                "w:b"
            )
        )

    run_xml.append(
        r_pr
    )

    texto = OxmlElement(
        "w:t"
    )

    texto.text = "1"

    run_xml.append(
        texto
    )

    fld_simple.append(
        run_xml
    )

    parrafo._p.append(
        fld_simple
    )


def configurar_documento(
    documento,
    distribucion,
    margenes_word
):
    paginas = distribucion.get(
        "paginas",
        []
    )

    if not paginas:
        raise RuntimeError(
            "La distribución no contiene páginas."
        )

    pagina = paginas[0]

    section = documento.sections[
        0
    ]

    section.page_width = Pt(
        numero(
            pagina.get(
                "ancho"
            ),
            612.0
        )
    )

    section.page_height = Pt(
        numero(
            pagina.get(
                "alto"
            ),
            792.0
        )
    )

    section.left_margin = Pt(
        margenes_word[
            "izquierdo"
        ]
    )

    section.right_margin = Pt(
        margenes_word[
            "derecho"
        ]
    )

    section.top_margin = Pt(
        margenes_word[
            "superior"
        ]
    )

    section.bottom_margin = Pt(
        margenes_word[
            "inferior"
        ]
    )

    configurar_borde_pagina(
        section,
        pagina.get(
            "borde_pagina",
            {}
        )
    )

    normal = documento.styles[
        "Normal"
    ]

    normal.font.name = "Aptos"
    normal.font.size = Pt(
        11
    )

    normal.paragraph_format.space_before = Pt(
        0
    )

    normal.paragraph_format.space_after = Pt(
        0
    )


# ================================================================
# UTILIDADES DE TABLAS
# ================================================================

def configurar_indentacion_tabla(
    tabla,
    puntos
):
    tbl_pr = tabla._tbl.tblPr

    tbl_ind = tbl_pr.find(
        qn(
            "w:tblInd"
        )
    )

    if tbl_ind is None:
        tbl_ind = OxmlElement(
            "w:tblInd"
        )

        tbl_pr.append(
            tbl_ind
        )

    tbl_ind.set(
        qn(
            "w:w"
        ),
        str(
            round(
                numero(
                    puntos
                )
                * 20
            )
        )
    )

    tbl_ind.set(
        qn(
            "w:type"
        ),
        "dxa"
    )


def configurar_ancho_celda(
    celda,
    puntos
):
    tc_pr = (
        celda
        ._tc
        .get_or_add_tcPr()
    )

    tc_w = tc_pr.find(
        qn(
            "w:tcW"
        )
    )

    if tc_w is None:
        tc_w = OxmlElement(
            "w:tcW"
        )

        tc_pr.append(
            tc_w
        )

    tc_w.set(
        qn(
            "w:w"
        ),
        str(
            round(
                numero(
                    puntos
                )
                * 20
            )
        )
    )

    tc_w.set(
        qn(
            "w:type"
        ),
        "dxa"
    )


def configurar_margenes_celda(
    celda,
    puntos=2.8
):
    tc_pr = (
        celda
        ._tc
        .get_or_add_tcPr()
    )

    tc_mar = tc_pr.find(
        qn(
            "w:tcMar"
        )
    )

    if tc_mar is None:
        tc_mar = OxmlElement(
            "w:tcMar"
        )

        tc_pr.append(
            tc_mar
        )

    valor = str(
        round(
            numero(
                puntos
            )
            * 20
        )
    )

    for lado in (
        "top",
        "left",
        "bottom",
        "right"
    ):
        margen = tc_mar.find(
            qn(
                f"w:{lado}"
            )
        )

        if margen is None:
            margen = OxmlElement(
                f"w:{lado}"
            )

            tc_mar.append(
                margen
            )

        margen.set(
            qn(
                "w:w"
            ),
            valor
        )

        margen.set(
            qn(
                "w:type"
            ),
            "dxa"
        )


# ================================================================
# ENCABEZADO Y PIE
# ================================================================

def configurar_encabezado_pie(
    documento,
    distribucion,
    margenes_word,
    mapa_efectos,
    ruta_textura,
    estadisticas
):
    paginas = distribucion.get(
        "paginas",
        []
    )

    if len(
        paginas
    ) < 2:
        return

    section = documento.sections[
        0
    ]

    pagina_1 = paginas[
        0
    ]

    pagina_ref = paginas[
        1
    ]

    ancho = numero(
        pagina_1.get(
            "ancho"
        ),
        612.0
    )

    alto = numero(
        pagina_1.get(
            "alto"
        ),
        792.0
    )

    section.different_first_page_header_footer = True
    section.header_distance = Pt(
        36.0
    )

    section.footer_distance = Pt(
        36.0
    )

    configurar_numero_pagina_inicio(
        section,
        0
    )

    if ruta_textura:
        agregar_fondo_mosaico_vml(
            section.first_page_header,
            ruta_textura,
            ancho,
            alto
        )

        agregar_fondo_mosaico_vml(
            section.header,
            ruta_textura,
            ancho,
            alto
        )

    p_header = section.header.paragraphs[
        0
    ]

    p_header.paragraph_format.space_before = Pt(
        0
    )

    p_header.paragraph_format.space_after = Pt(
        0
    )

    p_header.paragraph_format.line_spacing = 1.0

    p_header.paragraph_format.line_spacing_rule = (
        WD_LINE_SPACING.SINGLE
    )

    p_header.alignment = (
        WD_ALIGN_PARAGRAPH.CENTER
    )

    desactivar_ajuste_cuadricula(
        p_header
    )

    encabezado = pagina_ref.get(
        "encabezado"
    )

    if encabezado:
        efecto = obtener_efecto(
            mapa_efectos,
            pagina_ref.get(
                "numero",
                2
            ),
            identificador="encabezado",
            texto=encabezado.get(
                "texto"
            )
        )

        if efecto:
            run = p_header.add_run(
                str(
                    encabezado.get(
                        "texto",
                        ""
                    )
                )
            )

            estilo = dict(
                encabezado.get(
                    "estilo",
                    {}
                )
            )

            estilo[
                "tamano"
            ] = tamano_decorativo(
                encabezado
            )

            aplicar_formato_run(
                run,
                {},
                estilo
            )

            aplicar_efecto_run(
                run,
                efecto,
                estadisticas
            )

        else:
            agregar_lineas_formateadas(
                p_header,
                encabezado.get(
                    "lineas_origen",
                    []
                ),
                encabezado.get(
                    "estilo",
                    {}
                )
            )

    footer = section.footer

    p_footer_vacio = footer.paragraphs[
        0
    ]

    p_footer_vacio.clear()

    p_footer_vacio.paragraph_format.space_before = Pt(
        0
    )

    p_footer_vacio.paragraph_format.space_after = Pt(
        0
    )

    p_footer_vacio.paragraph_format.line_spacing = Pt(
        1
    )

    p_footer_vacio.paragraph_format.line_spacing_rule = (
        WD_LINE_SPACING.EXACTLY
    )

    pie = (
        pagina_ref.get(
            "pie"
        )
        or {}
    )

    texto_fecha = ""

    for zona in pie.get(
        "zonas",
        []
    ):
        if (
            zona.get(
                "posicion_horizontal"
            )
            == "izquierda"
        ):
            texto_fecha = str(
                zona.get(
                    "texto",
                    ""
                )
            )

            break

    ancho_footer = (
        ancho
        - margenes_word[
            "izquierdo"
        ]
        - margenes_word[
            "derecho"
        ]
    )

    tabla_footer = footer.add_table(
        rows=1,
        cols=2,
        width=Pt(
            ancho_footer
        )
    )

    tabla_footer.autofit = False
    tabla_footer.alignment = (
        WD_TABLE_ALIGNMENT.LEFT
    )

    configurar_indentacion_tabla(
        tabla_footer,
        0.0
    )

    tbl_pr = tabla_footer._tbl.tblPr

    bordes = OxmlElement(
        "w:tblBorders"
    )

    for nombre in (
        "top",
        "left",
        "bottom",
        "right",
        "insideH",
        "insideV"
    ):
        borde = OxmlElement(
            f"w:{nombre}"
        )

        borde.set(
            qn(
                "w:val"
            ),
            "nil"
        )

        bordes.append(
            borde
        )

    tbl_pr.append(
        bordes
    )

    ancho_columna = (
        ancho_footer
        / 2.0
    )

    for celda in tabla_footer.rows[
        0
    ].cells:
        configurar_ancho_celda(
            celda,
            ancho_columna
        )

        configurar_margenes_celda(
            celda,
            0
        )

    p_izq = tabla_footer.cell(
        0,
        0
    ).paragraphs[0]

    p_izq.clear()

    p_izq.alignment = (
        WD_ALIGN_PARAGRAPH.LEFT
    )

    p_izq.paragraph_format.space_before = Pt(
        0
    )

    p_izq.paragraph_format.space_after = Pt(
        0
    )

    run = p_izq.add_run(
        texto_fecha
    )

    run.font.name = "Aptos"
    run.font.size = Pt(
        12
    )

    p_der = tabla_footer.cell(
        0,
        1
    ).paragraphs[0]

    p_der.clear()

    p_der.alignment = (
        WD_ALIGN_PARAGRAPH.RIGHT
    )

    p_der.paragraph_format.space_before = Pt(
        0
    )

    p_der.paragraph_format.space_after = Pt(
        0
    )

    run = p_der.add_run(
        "Página "
    )

    run.font.name = "Aptos"
    run.font.size = Pt(
        12
    )

    agregar_campo_page(
        p_der,
        negrita=True
    )

    run = p_der.add_run(
        " de "
    )

    run.font.name = "Aptos"
    run.font.size = Pt(
        12
    )

    total = max(
        1,
        int(
            distribucion.get(
                "cantidad_paginas",
                len(
                    paginas
                )
            )
        )
        - 1
    )

    run = p_der.add_run(
        str(
            total
        )
    )

    run.font.name = "Aptos"
    run.font.size = Pt(
        12
    )

    run.bold = True


# ================================================================
# GEOMETRIA Y ESPACIADO
# ================================================================

def obtener_alineacion_elemento(
    elemento,
    tipo_pagina
):
    if tipo_pagina == "portada":
        lineas = elemento.get(
            "lineas_origen",
            []
        )

        if lineas:
            original = lineas[
                0
            ].get(
                "alineacion_estimada"
            )

            if original in (
                "izquierda",
                "centro",
                "derecha",
                "justificada"
            ):
                return original

    if (
        elemento.get(
            "tipo"
        )
        == "titulo"
        and tipo_pagina
        != "portada"
    ):
        return "izquierda"

    return elemento.get(
        "alineacion",
        "izquierda"
    )


def cuantizar_espacio_pdf(valor):
    valor = max(
        0.0,
        numero(
            valor
        )
    )

    if valor < 4.5:
        return 0.0

    if valor < 22.0:
        return 6.0

    if valor < 34.0:
        return 12.0

    if valor < 48.0:
        return 18.0

    return min(
        valor,
        72.0
    )


def inferir_espaciado_parrafo(
    elemento,
    tipo_pagina,
    margenes_word,
    es_primero
):
    caja = elemento.get(
        "caja",
        {}
    )

    top = numero(
        caja.get(
            "top"
        )
    )

    raw_antes = max(
        0.0,
        numero(
            elemento.get(
                "espacio_antes"
            )
        )
    )

    tipo = elemento.get(
        "tipo"
    )

    if tipo_pagina == "portada":
        if es_primero:
            antes = max(
                0.0,
                (
                    top
                    - margenes_word[
                        "superior"
                    ]
                )
            )

        else:
            antes = raw_antes

        return (
            antes,
            0.0
        )

    if es_primero:
        antes = max(
            0.0,
            (
                top
                - margenes_word[
                    "superior"
                ]
            )
        )

        return (
            antes,
            0.0
        )

    if tipo in {
        "entrada_indice",
        "texto_indice"
    }:
        return (
            0.0,
            0.0
        )

    if tipo == "titulo_indice":
        return (
            cuantizar_espacio_pdf(
                raw_antes
            ),
            0.0
        )

    if tipo == "titulo":
        nivel = int(
            numero(
                elemento.get(
                    "nivel"
                ),
                1
            )
        )

        if nivel <= 1:
            if raw_antes < 8.0:
                return (
                    0.0,
                    0.0
                )

            return (
                min(
                    18.0,
                    cuantizar_espacio_pdf(
                        raw_antes
                    )
                ),
                0.0
            )

        return (
            (
                0.0
                if raw_antes < 10.0
                else 6.0
            ),
            0.0
        )

    if tipo == "elemento_lista_titulo":
        return (
            (
                6.0
                if raw_antes >= 4.5
                else 0.0
            ),
            0.0
        )

    if tipo in {
        "elemento_lista_numerada",
        "elemento_lista_vineta"
    }:
        return (
            0.0,
            0.0
        )

    if tipo == "referencia":
        return (
            (
                6.0
                if raw_antes >= 4.5
                else 0.0
            ),
            0.0
        )

    if tipo == "parrafo":
        return (
            cuantizar_espacio_pdf(
                raw_antes
            ),
            0.0
        )

    return (
        cuantizar_espacio_pdf(
            raw_antes
        ),
        0.0
    )


def configurar_geometria_parrafo(
    parrafo,
    elemento,
    tipo_pagina,
    margenes_word,
    es_primero,
    espaciado_refinado=None,
    estadisticas=None,
    salto_pagina_antes=False
):
    formato = parrafo.paragraph_format

    caja = elemento.get(
        "caja",
        {}
    )

    x0 = numero(
        caja.get(
            "x0"
        )
    )

    margen_izq = margenes_word[
        "izquierdo"
    ]

    alineacion = obtener_alineacion_elemento(
        elemento,
        tipo_pagina
    )

    parrafo.alignment = convertir_alineacion(
        alineacion
    )

    tipo = elemento.get(
        "tipo"
    )

    sangrias = elemento.get(
        "sangrias",
        {}
    )

    francesa = max(
        0.0,
        numero(
            sangrias.get(
                "francesa"
            )
        )
    )

    if (
        tipo == "referencia"
        and francesa > 0.5
    ):
        base_primera_linea = max(
            0.0,
            (
                x0
                - margen_izq
            )
        )

        formato.left_indent = Pt(
            base_primera_linea
            + francesa
        )

        formato.first_line_indent = Pt(
            -francesa
        )

    elif (
        tipo_pagina == "portada"
        and alineacion == "centro"
    ):
        formato.left_indent = Pt(
            0
        )

    else:
        formato.left_indent = Pt(
            x0
            - margen_izq
        )

        primera = numero(
            sangrias.get(
                "primera_linea"
            )
        )

        if abs(
            primera
        ) > 0.2:
            formato.first_line_indent = Pt(
                primera
            )

    antes, despues, uso_refinado = resolver_espaciado_word(
        espaciado_refinado,
        elemento,
        tipo_pagina,
        margenes_word,
        es_primero
    )

    if (
        uso_refinado
        and estadisticas is not None
    ):
        estadisticas[
            "espaciados_refinados_aplicados"
        ] += 1

    formato.space_before = Pt(
        antes
    )

    formato.space_after = Pt(
        despues
    )

    if tipo_pagina == "portada":
        tamano = max(
            1.0,
            tamano_decorativo(
                elemento
            )
        )

        formato.line_spacing = Pt(
            tamano
        )

        formato.line_spacing_rule = (
            WD_LINE_SPACING.EXACTLY
        )

    elif tipo == "titulo_indice":
        formato.line_spacing = 1.5

        formato.line_spacing_rule = (
            WD_LINE_SPACING.SINGLE
        )

    else:
        formato.line_spacing = 1.0

        formato.line_spacing_rule = (
            WD_LINE_SPACING.SINGLE
        )

    formato.page_break_before = bool(
        salto_pagina_antes
    )

    desactivar_ajuste_cuadricula(
        parrafo
    )


# ================================================================
# INDICE
# ================================================================

def agregar_tab_stop_underscore(
    parrafo,
    posicion_pt
):
    p_pr = (
        parrafo
        ._p
        .get_or_add_pPr()
    )

    tabs = p_pr.find(
        qn(
            "w:tabs"
        )
    )

    if tabs is None:
        tabs = OxmlElement(
            "w:tabs"
        )

        p_pr.append(
            tabs
        )

    tab = OxmlElement(
        "w:tab"
    )

    tab.set(
        qn(
            "w:val"
        ),
        "right"
    )

    tab.set(
        qn(
            "w:leader"
        ),
        "underscore"
    )

    tab.set(
        qn(
            "w:pos"
        ),
        str(
            int(
                round(
                    posicion_pt
                    * 20.0
                )
            )
        )
    )

    tabs.append(
        tab
    )


def agregar_entrada_indice(
    documento,
    elemento,
    margenes_word,
    ancho_pagina,
    es_primero=False,
    espaciado_refinado=None,
    estadisticas=None,
    salto_pagina_antes=False
):
    parrafo = documento.add_paragraph()

    formato = parrafo.paragraph_format

    caja = elemento.get(
        "caja",
        {}
    )

    x0 = numero(
        caja.get(
            "x0"
        )
    )

    formato.left_indent = Pt(
        x0
        - margenes_word[
            "izquierdo"
        ]
    )

    antes, despues, uso_refinado = resolver_espaciado_word(
        espaciado_refinado,
        elemento,
        "indice",
        margenes_word,
        es_primero
    )

    if (
        uso_refinado
        and estadisticas is not None
    ):
        estadisticas[
            "espaciados_refinados_aplicados"
        ] += 1

    formato.space_before = Pt(
        antes
    )

    formato.space_after = Pt(
        despues
    )

    formato.line_spacing = 1.0
    formato.line_spacing_rule = (
        WD_LINE_SPACING.SINGLE
    )

    formato.page_break_before = bool(
        salto_pagina_antes
    )

    parrafo.alignment = (
        WD_ALIGN_PARAGRAPH.LEFT
    )

    desactivar_ajuste_cuadricula(
        parrafo
    )

    ancho_util = (
        ancho_pagina
        - margenes_word[
            "izquierdo"
        ]
        - margenes_word[
            "derecho"
        ]
    )

    posicion_tab = ancho_util

    agregar_tab_stop_underscore(
        parrafo,
        posicion_tab
    )

    estilo = elemento.get(
        "estilo",
        {}
    )

    texto_limpio = limpiar_texto_indice(
        elemento.get(
            "texto",
            ""
        )
    )

    run_texto = parrafo.add_run(
        texto_limpio
    )

    aplicar_formato_run(
        run_texto,
        {},
        estilo
    )

    parrafo.add_run(
        "\t"
    )

    run_pagina = parrafo.add_run(
        obtener_pagina_indice(
            elemento
        )
    )

    aplicar_formato_run(
        run_pagina,
        {},
        estilo
    )

    return parrafo


def agregar_elemento_texto(
    documento,
    elemento,
    tipo_pagina,
    margenes_word,
    es_primero,
    efecto,
    espaciado_refinado,
    estadisticas,
    salto_pagina_antes=False
):
    parrafo = documento.add_paragraph()

    configurar_geometria_parrafo(
        parrafo,
        elemento,
        tipo_pagina,
        margenes_word,
        es_primero,
        espaciado_refinado=espaciado_refinado,
        estadisticas=estadisticas,
        salto_pagina_antes=salto_pagina_antes
    )

    reflejo = (
        (efecto or {})
        .get(
            "reflejo_final",
            {}
        )
    )

    relleno = (
        (efecto or {})
        .get(
            "relleno_final",
            {}
        )
    )

    usa_run_unico = (
        bool(
            reflejo.get(
                "presente"
            )
        )
        or relleno.get(
            "tipo"
        )
        == "degradado"
    )

    if usa_run_unico:
        run = parrafo.add_run(
            str(
                elemento.get(
                    "texto",
                    ""
                )
            )
        )

        estilo = dict(
            elemento.get(
                "estilo",
                {}
            )
        )

        estilo[
            "tamano"
        ] = tamano_decorativo(
            elemento
        )

        aplicar_formato_run(
            run,
            {},
            estilo
        )

        aplicar_efecto_run(
            run,
            efecto,
            estadisticas
        )

        return parrafo

    lineas = elemento.get(
        "lineas_origen",
        []
    )

    indice_inicio_runs = len(
        parrafo.runs
    )

    if lineas:
        agregar_lineas_formateadas(
            parrafo,
            lineas,
            elemento.get(
                "estilo",
                {}
            )
        )

    else:
        run = parrafo.add_run(
            str(
                elemento.get(
                    "texto",
                    ""
                )
            )
        )

        aplicar_formato_run(
            run,
            {},
            elemento.get(
                "estilo",
                {}
            )
        )

    if efecto:
        color_final = efecto.get(
            "color_final",
            {}
        )

        if (
            color_final.get(
                "origen"
            )
            == "analisis_visual"
        ):
            for run in parrafo.runs[
                indice_inicio_runs:
            ]:
                aplicar_efecto_run(
                    run,
                    efecto,
                    estadisticas
                )

    return parrafo


# ================================================================
# TABLAS EDITABLES
# ================================================================

def configurar_bordes_tabla(
    tabla,
    grosor
):
    tbl_pr = tabla._tbl.tblPr

    bordes = tbl_pr.find(
        qn(
            "w:tblBorders"
        )
    )

    if bordes is None:
        bordes = OxmlElement(
            "w:tblBorders"
        )

        tbl_pr.append(
            bordes
        )

    sz = str(
        max(
            2,
            round(
                numero(
                    grosor,
                    0.4
                )
                * 8
            )
        )
    )

    for nombre in (
        "top",
        "left",
        "bottom",
        "right",
        "insideH",
        "insideV"
    ):
        borde = bordes.find(
            qn(
                f"w:{nombre}"
            )
        )

        if borde is None:
            borde = OxmlElement(
                f"w:{nombre}"
            )

            bordes.append(
                borde
            )

        borde.set(
            qn(
                "w:val"
            ),
            "single"
        )

        borde.set(
            qn(
                "w:sz"
            ),
            sz
        )

        borde.set(
            qn(
                "w:space"
            ),
            "0"
        )

        borde.set(
            qn(
                "w:color"
            ),
            "000000"
        )


def configurar_relleno_celda(
    celda,
    rgb
):
    if (
        not rgb
        or len(
            rgb
        ) < 3
    ):
        return

    tc_pr = (
        celda
        ._tc
        .get_or_add_tcPr()
    )

    shd = tc_pr.find(
        qn(
            "w:shd"
        )
    )

    if shd is None:
        shd = OxmlElement(
            "w:shd"
        )

        tc_pr.append(
            shd
        )

    color = "".join(
        f"{max(0, min(255, int(valor))):02X}"
        for valor in rgb[:3]
    )

    shd.set(
        qn(
            "w:val"
        ),
        "clear"
    )

    shd.set(
        qn(
            "w:color"
        ),
        "auto"
    )

    shd.set(
        qn(
            "w:fill"
        ),
        color
    )


def agregar_contenido_celda(
    celda_word,
    datos_celda
):
    parrafo = celda_word.paragraphs[
        0
    ]

    parrafo.clear()

    parrafo.alignment = convertir_alineacion(
        datos_celda.get(
            "alineacion"
        )
    )

    parrafo.paragraph_format.space_before = Pt(
        0
    )

    parrafo.paragraph_format.space_after = Pt(
        0
    )

    parrafo.paragraph_format.line_spacing = 1.0
    parrafo.paragraph_format.line_spacing_rule = (
        WD_LINE_SPACING.SINGLE
    )

    desactivar_ajuste_cuadricula(
        parrafo
    )

    lineas = [
        {
            "texto": linea.get(
                "texto",
                ""
            ),
            "palabras": linea.get(
                "palabras",
                []
            )
        }
        for linea in datos_celda.get(
            "lineas_visuales",
            []
        )
    ]

    if lineas:
        agregar_lineas_formateadas(
            parrafo,
            lineas,
            datos_celda.get(
                "estilo",
                {}
            )
        )

    else:
        run = parrafo.add_run(
            str(
                datos_celda.get(
                    "texto",
                    ""
                )
            )
        )

        aplicar_formato_run(
            run,
            {},
            datos_celda.get(
                "estilo",
                {}
            )
        )


def agregar_tabla(
    documento,
    elemento,
    margenes_word,
    espaciado_refinado=None,
    estadisticas=None,
    salto_pagina_antes=False,
    espacio_antes_en_ancla=True
):
    datos = elemento.get(
        "datos_tabla",
        {}
    )

    filas = datos.get(
        "filas_reconstruidas",
        []
    )

    columnas = int(
        datos.get(
            "columnas",
            0
        )
    )

    if (
        not filas
        or columnas <= 0
    ):
        return None

    if salto_pagina_antes:
        ancla = documento.add_paragraph()

        ancla.paragraph_format.page_break_before = True
        ancla.paragraph_format.space_before = Pt(
            0
        )

        espacio_antes = 0.0

        if (
            espacio_antes_en_ancla
            and isinstance(
                espaciado_refinado,
                dict
            )
        ):
            espacio_antes = max(
                0.0,
                numero(
                    espaciado_refinado.get(
                        "espacio_word_antes"
                    )
                )
            )

            if estadisticas is not None:
                estadisticas[
                    "espaciados_refinados_aplicados"
                ] += 1

        ancla.paragraph_format.space_after = Pt(
            espacio_antes
        )

        ancla.paragraph_format.line_spacing = Pt(
            1
        )

        ancla.paragraph_format.line_spacing_rule = (
            WD_LINE_SPACING.EXACTLY
        )

    tabla = documento.add_table(
        rows=len(
            filas
        ),
        cols=columnas
    )

    tabla.autofit = False
    tabla.alignment = (
        WD_TABLE_ALIGNMENT.LEFT
    )

    limites_x = datos.get(
        "limites_x",
        []
    )

    if limites_x:
        configurar_indentacion_tabla(
            tabla,
            (
                numero(
                    limites_x[0]
                )
                - margenes_word[
                    "izquierdo"
                ]
            )
        )

    configurar_bordes_tabla(
        tabla,
        datos.get(
            "grosor_borde_estimado",
            0.4
        )
    )

    anchos = datos.get(
        "anchos_columnas",
        []
    )

    for indice, ancho in enumerate(
        anchos
    ):
        if indice < len(
            tabla.columns
        ):
            tabla.columns[
                indice
            ].width = Pt(
                numero(
                    ancho
                )
            )

    for i, datos_fila in enumerate(
        filas
    ):
        fila = tabla.rows[
            i
        ]

        fila.height = Pt(
            numero(
                datos_fila.get(
                    "alto"
                ),
                12.0
            )
        )

        fila.height_rule = (
            WD_ROW_HEIGHT_RULE.EXACTLY
        )

        for j, datos_celda in enumerate(
            datos_fila.get(
                "celdas",
                []
            )
        ):
            if j >= columnas:
                continue

            celda = fila.cells[
                j
            ]

            if j < len(
                anchos
            ):
                configurar_ancho_celda(
                    celda,
                    anchos[j]
                )

            configurar_margenes_celda(
                celda,
                2.8
            )

            configurar_relleno_celda(
                celda,
                datos_celda.get(
                    "relleno_rgb"
                )
            )

            vertical = str(
                datos_celda.get(
                    "alineacion_vertical",
                    "centro"
                )
            )

            if vertical == "arriba":
                celda.vertical_alignment = (
                    WD_CELL_VERTICAL_ALIGNMENT.TOP
                )

            elif vertical == "abajo":
                celda.vertical_alignment = (
                    WD_CELL_VERTICAL_ALIGNMENT.BOTTOM
                )

            else:
                celda.vertical_alignment = (
                    WD_CELL_VERTICAL_ALIGNMENT.CENTER
                )

            agregar_contenido_celda(
                celda,
                datos_celda
            )

    return tabla


# ================================================================
# GRAFICOS DE PORTADA
# ================================================================

def imagen_debe_insertarse(
    imagen,
    nombres_reflejo
):
    nombre = str(
        imagen.get(
            "nombre"
        )
        or ""
    )

    rol = str(
        imagen.get(
            "rol_estimado"
        )
        or ""
    ).lower()

    if nombre in nombres_reflejo:
        return False

    if rol == "efecto_grafico_texto":
        return False

    return True


def agregar_imagenes_portada(
    pagina,
    parrafo_ancla,
    mapa_efectos
):
    numero_pagina = int(
        numero(
            pagina.get(
                "numero"
            ),
            1
        )
    )

    reflejos = nombres_imagenes_reflejo(
        mapa_efectos,
        numero_pagina
    )

    for imagen in (
        pagina
        .get(
            "recursos_graficos",
            {}
        )
        .get(
            "imagenes",
            []
        )
    ):
        if not imagen_debe_insertarse(
            imagen,
            reflejos
        ):
            continue

        caja = imagen.get(
            "caja",
            {}
        )

        agregar_imagen_flotante(
            parrafo_ancla,
            imagen.get(
                "archivo_extraido"
            ),
            numero(
                caja.get(
                    "x0"
                )
            ),
            numero(
                caja.get(
                    "top"
                )
            ),
            (
                numero(
                    caja.get(
                        "x1"
                    )
                )
                - numero(
                    caja.get(
                        "x0"
                    )
                )
            ),
            (
                numero(
                    caja.get(
                        "bottom"
                    )
                )
                - numero(
                    caja.get(
                        "top"
                    )
                )
            ),
            True
        )


# ================================================================
# GENERACION DE PAGINAS
# ================================================================

def agregar_pagina(
    documento,
    pagina,
    margenes_word,
    numero_pagina,
    mapa_efectos,
    mapa_espaciados,
    ancho_pagina,
    estadisticas
):
    elementos = pagina.get(
        "elementos",
        []
    )

    primer_elemento_generado = False
    ultimo_parrafo = None

    for indice, elemento in enumerate(
        elementos,
        start=1
    ):
        tipo = elemento.get(
            "tipo"
        )

        identificador = (
            f"elemento_{indice:03d}"
        )

        efecto = obtener_efecto(
            mapa_efectos,
            numero_pagina,
            identificador=identificador,
            texto=elemento.get(
                "texto"
            )
        )

        espaciado_refinado = obtener_espaciado_refinado(
            mapa_espaciados,
            numero_pagina,
            indice_elemento=(
                indice
                - 1
            ),
            tipo=tipo,
            texto=elemento.get(
                "texto"
            )
        )

        es_primero = (
            not primer_elemento_generado
        )

        salto_pagina_antes = (
            numero_pagina > 1
            and es_primero
        )

        if tipo == "tabla":
            espacio_tabla = 0.0

            if isinstance(
                espaciado_refinado,
                dict
            ):
                espacio_tabla = max(
                    0.0,
                    numero(
                        espaciado_refinado.get(
                            "espacio_word_antes"
                        )
                    )
                )

            if (
                ultimo_parrafo is not None
                and espacio_tabla > 0.0
            ):
                ultimo_parrafo.paragraph_format.space_after = Pt(
                    espacio_tabla
                )

                estadisticas[
                    "espaciados_refinados_aplicados"
                ] += 1

            agregar_tabla(
                documento,
                elemento,
                margenes_word,
                espaciado_refinado=espaciado_refinado,
                estadisticas=estadisticas,
                salto_pagina_antes=salto_pagina_antes,
                espacio_antes_en_ancla=(
                    ultimo_parrafo is None
                )
            )

            primer_elemento_generado = True
            ultimo_parrafo = None
            continue

        if tipo == "entrada_indice":
            parrafo = agregar_entrada_indice(
                documento,
                elemento,
                margenes_word,
                ancho_pagina,
                es_primero=es_primero,
                espaciado_refinado=espaciado_refinado,
                estadisticas=estadisticas,
                salto_pagina_antes=salto_pagina_antes
            )

        else:
            parrafo = agregar_elemento_texto(
                documento,
                elemento,
                pagina.get(
                    "tipo"
                ),
                margenes_word,
                es_primero,
                efecto,
                espaciado_refinado,
                estadisticas,
                salto_pagina_antes=salto_pagina_antes
            )

        if (
            numero_pagina == 1
            and es_primero
        ):
            agregar_imagenes_portada(
                pagina,
                parrafo,
                mapa_efectos
            )

        primer_elemento_generado = True
        ultimo_parrafo = parrafo


# ================================================================
# VALIDACION DEL DOCX
# ================================================================

def inspeccionar_efectos_guardados(
    ruta_docx
):
    if not ruta_docx.is_file():
        return {
            "text_fill": 0,
            "reflection": 0
        }

    with zipfile.ZipFile(
        ruta_docx,
        "r"
    ) as zip_docx:
        total_text_fill = 0
        total_reflection = 0

        for nombre in zip_docx.namelist():
            if not (
                nombre.startswith(
                    "word/"
                )
                and nombre.endswith(
                    ".xml"
                )
            ):
                continue

            datos = zip_docx.read(
                nombre
            )

            total_text_fill += datos.count(
                b"<w14:textFill"
            )

            total_reflection += datos.count(
                b"<w14:reflection"
            )

    return {
        "text_fill": total_text_fill,
        "reflection": total_reflection
    }


def imprimir_diagnostico(
    distribucion,
    margenes_word,
    mapa_efectos,
    mapa_espaciados,
    ruta_textura
):
    print()

    print(
        "Configuración reconstruida:"
    )

    print(
        (
            "  Márgenes Word inferidos: "
            f"izq={margenes_word['izquierdo']:.2f} pt | "
            f"der={margenes_word['derecho']:.2f} pt | "
            f"sup={margenes_word['superior']:.2f} pt | "
            f"inf={margenes_word['inferior']:.2f} pt"
        )
    )

    print(
        (
            "  Textura mosaico: "
            f"{'sí' if ruta_textura else 'no'}"
        )
    )

    print(
        (
            "  Efectos refinados cargados: "
            f"{len(mapa_efectos['por_id'])}"
        )
    )

    print(
        (
            "  Espaciados refinados cargados: "
            f"{len(mapa_espaciados.get('por_indice', {}))}"
        )
    )

    print(
        (
            "  Páginas fuente: "
            f"{len(distribucion.get('paginas', []))}"
        )
    )

    print(
        "  Saltos de página: pageBreakBefore en el primer elemento de cada página fuente"
    )

    print(
        "  Espaciado: usa espaciados_pdf_refinados.json; si falta una entrada, aplica la heurística interna"
    )

    print()


# ================================================================
# EJECUCION PRINCIPAL
# ================================================================

def main():
    distribucion = cargar_json(
        RUTA_DISTRIBUCION
    )

    datos_efectos = cargar_json_opcional(
        RUTA_EFECTOS_REFINADOS
    )

    mapa_efectos = construir_mapa_efectos(
        datos_efectos
    )

    datos_espaciados = cargar_json_opcional(
        RUTA_ESPACIADOS_REFINADOS
    )

    mapa_espaciados = construir_mapa_espaciados(
        datos_espaciados
    )

    paginas = distribucion.get(
        "paginas",
        []
    )

    if not paginas:
        raise RuntimeError(
            "No hay páginas para generar el DOCX."
        )

    CARPETA_SALIDA.mkdir(
        parents=True,
        exist_ok=True
    )

    margenes_word = inferir_margenes_word(
        distribucion
    )

    ruta_textura = obtener_textura_fondo(
        distribucion
    )

    imprimir_diagnostico(
        distribucion,
        margenes_word,
        mapa_efectos,
        mapa_espaciados,
        ruta_textura
    )

    estadisticas = {
        "efectos_aplicados": 0,
        "text_fill_aplicados": 0,
        "reflejos_aplicados": 0,
        "espaciados_refinados_aplicados": 0
    }

    print(
        "Generando DOCX editable..."
    )

    print()

    documento = Document()

    configurar_documento(
        documento,
        distribucion,
        margenes_word
    )

    configurar_encabezado_pie(
        documento,
        distribucion,
        margenes_word,
        mapa_efectos,
        ruta_textura,
        estadisticas
    )

    ancho_pagina = numero(
        paginas[0].get(
            "ancho"
        ),
        612.0
    )

    for pagina in paginas:
        numero_pagina = int(
            numero(
                pagina.get(
                    "numero"
                ),
                0
            )
        )

        print(
            f"Página {numero_pagina}..."
        )

        agregar_pagina(
            documento,
            pagina,
            margenes_word,
            numero_pagina,
            mapa_efectos,
            mapa_espaciados,
            ancho_pagina,
            estadisticas
        )

    documento.save(
        RUTA_DOCX
    )

    guardados = inspeccionar_efectos_guardados(
        RUTA_DOCX
    )

    print()

    print(
        "DOCX generado."
    )

    print(
        f"Salida: {RUTA_DOCX}"
    )

    print()

    print(
        "Diagnóstico de efectos:"
    )

    print(
        (
            "  Elementos a los que se intentó aplicar color/efecto: "
            f"{estadisticas['efectos_aplicados']}"
        )
    )

    print(
        (
            "  textFill insertados durante la generación: "
            f"{estadisticas['text_fill_aplicados']}"
        )
    )

    print(
        (
            "  reflection insertados durante la generación: "
            f"{estadisticas['reflejos_aplicados']}"
        )
    )

    print(
        (
            "  Espaciados refinados aplicados: "
            f"{estadisticas['espaciados_refinados_aplicados']}"
        )
    )

    print(
        (
            "  w14:textFill encontrados dentro del DOCX guardado: "
            f"{guardados['text_fill']}"
        )
    )

    print(
        (
            "  w14:reflection encontrados dentro del DOCX guardado: "
            f"{guardados['reflection']}"
        )
    )

    if (
        estadisticas[
            "text_fill_aplicados"
        ] > 0
        and guardados[
            "text_fill"
        ] == 0
    ):
        print()

        print(
            "ADVERTENCIA: se intentaron aplicar efectos, pero no quedaron guardados en el OOXML."
        )

    print()


if __name__ == "__main__":
    main()